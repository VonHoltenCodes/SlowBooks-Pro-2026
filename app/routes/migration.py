# ============================================================================
# Migration — one route pair for every "migrate from X" source.
#
# GET  /api/migration/sources               → picker metadata
# POST /api/migration/{source}/dry-run      → validate, write nothing
# POST /api/migration/{source}/import       → dry-run-gated import
#
# The Xero and MYOB-specific routes (/api/xero-import, /api/myob-import)
# remain as compatibility aliases; new sources exist only here.
# ============================================================================

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy.orm import Session

from app.database import get_db
from app.services import (
    gnucash_import,
    myob_import,
    qbo_export_import,
    sage_import,
    wave_import,
    xero_import,
    zoho_import,
)
from app.services.settings_service import set_setting
from app.services.upload_limits import read_limited

router = APIRouter(prefix="/api/migration", tags=["migration"])

# source key → (label, dialect module, chart_setup_source value)
SOURCES = {
    "xero": ("Xero", xero_import, "xero_import"),
    "myob": ("MYOB", myob_import, "myob_import"),
    "sage": ("Sage 50", sage_import, "sage_import"),
    "wave": ("Wave", wave_import, "wave_import"),
    "zoho": ("Zoho Books", zoho_import, "zoho_import"),
    "gnucash": ("GnuCash", gnucash_import, "gnucash_import"),
    # QuickBooks Online's own Export data bundle, no Intuit developer app
    # needed: the zip as downloaded, or its workbooks, or those saved as CSV.
    "qbo_export": ("QuickBooks Online export", qbo_export_import, "qbo_export_import"),
}


def _source(key: str):
    entry = SOURCES.get(key)
    if not entry:
        raise HTTPException(status_code=400, detail=f"Unknown migration source: {key}")
    return entry


def _sheet_as_csv(content: bytes, computed: bool = False) -> str:
    """An Excel workbook's first sheet as CSV text, cells as they read:
    QuickBooks Online exports its reports and lists as .xlsx, with every
    figure as formula text ("=23683.16") and labels indented by spaces, and
    both survive this pass for the dialect to read. With `computed`, a
    formula cell reads as the value Excel last computed for it, when the
    file carries one (QuickBooks Desktop's timesheet report writes its
    Total rows as =ROUND(SUM(...),5); 2.23.4 gate, NEW-132); a formula
    without a computed value stays as its text."""
    import csv
    import io
    from datetime import date, datetime, time, timedelta

    try:
        import openpyxl
    except ImportError:  # pragma: no cover - a build without the reader
        raise HTTPException(
            status_code=400,
            detail="This build cannot read .xlsx files; save the sheet as CSV and upload that.",
        )
    try:
        book = openpyxl.load_workbook(
            io.BytesIO(content), read_only=True, data_only=False
        )
    except Exception:
        raise HTTPException(
            status_code=400, detail="The .xlsx file could not be read as a workbook"
        )
    sheet = book.worksheets[0]
    computed_rows = None
    if computed:
        try:
            values = openpyxl.load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
            computed_rows = values.worksheets[0].iter_rows(values_only=True)
        except Exception:
            computed_rows = None
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    for row in sheet.iter_rows(values_only=True):
        cached = next(computed_rows, None) if computed_rows is not None else None
        if cached is not None:
            row = tuple(
                (
                    cached[i]
                    if isinstance(cell, str)
                    and cell.startswith("=")
                    and i < len(cached)
                    and cached[i] is not None
                    else cell
                )
                for i, cell in enumerate(row)
            )
        cells = []
        for cell in row:
            if cell is None:
                cells.append("")
            elif isinstance(cell, datetime):
                cells.append(cell.strftime("%m/%d/%Y"))
            elif isinstance(cell, date):
                cells.append(cell.strftime("%m/%d/%Y"))
            elif isinstance(cell, time):
                # a duration Excel took for a time of day ("0:20" in a
                # Minutes-format timesheet): hours and minutes (NEW-137)
                cells.append(f"{cell.hour}:{cell.minute:02d}")
            elif isinstance(cell, timedelta):
                # a duration past a day, as Excel keeps [h]:mm
                minutes = int(round(cell.total_seconds() / 60))
                cells.append(f"{minutes // 60}:{minutes % 60:02d}")
            elif isinstance(cell, float) and cell.is_integer():
                cells.append(str(int(cell)))
            else:
                cells.append(str(cell))
        writer.writerow(cells)
    book.close()
    return out.getvalue()


def _members(filename: str, content: bytes) -> list[tuple[str, bytes]]:
    """An upload as the files it holds: a zip's members (a vendor's bundle,
    whose unknown members are left alone), or the file itself."""
    import io
    import posixpath
    import zipfile

    if not (filename or "").lower().endswith(".zip"):
        return [(filename, content)]
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile:
        raise HTTPException(status_code=400, detail=f"{filename} is not a zip file")
    members = []
    with archive:
        for info in archive.infolist():
            name = posixpath.basename(info.filename)
            if (
                info.is_dir()
                or not name
                or name.startswith(".")
                or "__MACOSX" in info.filename
            ):
                continue
            members.append((name, archive.read(info)))
    return members


async def bundle_from_uploads(files: list[UploadFile], classify) -> dict:
    bundle: dict[str, str] = {}
    unrecognized = []
    for file in files:
        content = await read_limited(file, label=f"Migration file {file.filename}")
        zipped = (file.filename or "").lower().endswith(".zip")
        for name, data in _members(file.filename, content):
            kind = classify(name)
            if kind == "skip" or (not kind and zipped):
                continue  # a file the source says to leave alone, or a zip's stray
            if not kind:
                unrecognized.append(name)
                continue
            if name.lower().endswith(".xlsx"):
                text = _sheet_as_csv(data)
            else:
                try:
                    text = data.decode("utf-8-sig")
                except UnicodeDecodeError:
                    text = data.decode("latin-1")
            _add_to_bundle(bundle, kind, name, text)
    if unrecognized:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Could not classify files by name: {unrecognized}. Expected "
                f"filenames containing 'chart'/'accounts', "
                f"'ledger'/'journal'/'transactions', or 'trial'."
            ),
        )
    return bundle


def _add_to_bundle(bundle: dict, kind: str, name: str, text: str) -> None:
    if True:
        if kind in bundle:
            # Multiple files of one kind (MYOB exports journals per type):
            # concatenate, dropping the duplicate header row — which must
            # match the first file's, or the files aren't the same table.
            first_header = bundle[kind].split("\n", 1)[0].strip()
            header, _, body = text.partition("\n")
            if header.strip() != first_header:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"{name} has different columns than the other "
                        f"{kind} file(s); export them with the same settings"
                    ),
                )
            bundle[kind] = bundle[kind].rstrip("\n") + "\n" + body
        else:
            bundle[kind] = text


@router.get("/sources")
def list_sources():
    return [{"key": key, "label": label} for key, (label, _, _) in SOURCES.items()]


@router.post("/{source}/dry-run")
async def migration_dry_run(
    source: str, files: list[UploadFile] = File(...), db: Session = Depends(get_db)
):
    _, module, _ = _source(source)
    bundle = await bundle_from_uploads(files, module.classify_filename)
    return module.dry_run(db, bundle)


@router.post("/{source}/import")
async def migration_import(
    source: str, files: list[UploadFile] = File(...), db: Session = Depends(get_db)
):
    _, module, chart_source = _source(source)
    bundle = await bundle_from_uploads(files, module.classify_filename)
    result = module.run_import(db, bundle)
    if result["ok"]:
        from datetime import date

        set_setting(db, "chart_setup_source", chart_source)
        set_setting(db, "chart_setup_ready_at", date.today().isoformat())
        db.commit()
    return result
